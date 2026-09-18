# -*- coding: utf-8 -*-
"""工作汇报 md -> Word（一次性脚本）。图插在对应章节，产物与 md 同目录同名 .docx。"""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

MD = Path("docs/工作汇报-20260917-0918-聚类定档与纠错.md")
FIGDIR = Path("docs/工作汇报-20260917-0918-图")
OUT = MD.with_suffix(".docx")

doc = Document()
style = doc.styles["Normal"]
style.font.name = "Calibri"
style.font.size = Pt(11)
style.element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")


def add_runs(p, text):
    """把 **bold** 语法转成 run"""
    for i, seg in enumerate(re.split(r"\*\*", text)):
        if not seg:
            continue
        r = p.add_run(seg)
        r.bold = i % 2 == 1


def para(text, center=False):
    p = doc.add_paragraph()
    add_runs(p, text)
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def heading(text, level):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    return p


def add_table(rows):
    t = doc.add_table(rows=0, cols=len(rows[0]))
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for ci, val in enumerate(row):
            cells[ci].text = ""
            p = cells[ci].paragraphs[0]
            add_runs(p, val)
            if ri == 0:
                for r in p.runs:
                    r.bold = True


def shot(fname, caption):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(FIGDIR / fname), width=Cm(14.5))
    c = doc.add_paragraph()
    c.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = c.add_run(caption)
    r.font.size = Pt(9)
    r.italic = True


FIG1_AFTER = "结论：**归属原子性优先**"
FIG2_AFTER = "λ=1.0 会帮倒忙"

lines = MD.read_text(encoding="utf-8").splitlines()
i = 0
inserted = {1: False, 2: False, 3: False}
while i < len(lines):
    line = lines[i].rstrip()
    if not line.strip() or line.strip() == "---":
        i += 1
        continue
    if line.startswith("### "):
        heading(line[4:], 2)
    elif line.startswith("## "):
        heading(line[3:], 1)
    elif line.startswith("# "):
        t = doc.add_heading(line[2:], level=0)
        for r in t.runs:
            r._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), "黑体")
    elif line.startswith("|"):
        rows = []
        while i < len(lines) and lines[i].strip().startswith("|"):
            cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
            if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                rows.append(cells)
            i += 1
        add_table(rows)
        continue
    elif line.startswith("- "):
        para("• " + line[2:])
    else:
        para(line)
        if not inserted[1] and FIG1_AFTER in line:
            inserted[1] = True
            shot("fig1_VAD分割调试对比听感台.png",
                 "图 1 ｜ 说话人分离调试听感台（VAD 分割对比）：同一段无缝拼接压力音频，"
                 "逐段试听并按说话人着色，上方 GT 25 段 vs 下方系统 21 段逐段对照——"
                 "VAD 切句质量与聚类归属的可视化调试依据。")
        if not inserted[2] and FIG2_AFTER in line:
            inserted[2] = True
            shot("fig2_热词筛选审查台.png",
                 "图 2 ｜ 热词筛选审查台：2316 个候选热词按「会议口语高频（AISHELL-4 转写挖出，≥3 次）」"
                 "「英文术语（科技/办公）」等分组逐词人工审查，坚持「小而准、宁可删错不可滥放」——"
                 "避免裸替换类误纠（如「校长→晓庄」）的词表治理过程。")
    i += 1

# 附：降噪对比（明日计划第 3 条相关，图放在文末附加节）
heading("附：降噪方案对比试听（GTCRN / DeepFilterNet3 / FRCRN）", 1)
para("同一批录音样本分别过三种降噪模型的试听对比台：每条样本并列展示「原声（未降噪）/ GTCRN（西工大）/ "
     "DeepFilterNet3（德国）/ FRCRN（阿里）」四个版本波形，人工盲听听感并结合作管线实测"
     "（DER 19.0% → GTCRN 12.3% → FRCRN 9.1%，但 FRCRN 误伤人声且慢 30 倍）选定 GTCRN 上线。")
shot("fig3_降噪对比试听台.png",
     "图 3 ｜ 降噪对比试听台：同一样本四种版本（原声/GTCRN/DeepFilterNet3/FRCRN）并列试听，"
     "历史记录持久化，支撑降噪模型选型决策。")

doc.save(str(OUT))
print("saved:", OUT, "| figs inserted:", inserted)
